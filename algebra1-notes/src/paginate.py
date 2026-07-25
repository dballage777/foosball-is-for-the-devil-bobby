#!/usr/bin/env python3
"""Post-process the .docx to enforce professional pagination:
   - every paragraph: keep lines together (no single-line widows/orphans)
   - every table row: cantSplit (a row never breaks across a page)
   - whole tables stay together (keep-with-next chained through all but last row)
   - a heading/directions paragraph stays glued to the table that follows it
   - list items stay together as a block (keep-with-next on all but the last)
   These are exactly the paragraph settings Google Docs exposes as
   'Keep lines together' / 'Keep with next line' and the table 'cantSplit' flag.
"""
import sys
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = sys.argv[1]
doc = Document(path)


def has_numpr(p):
    pPr = p._p.pPr
    return pPr is not None and pPr.find(qn('w:numPr')) is not None


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def all_paragraphs(el, parent):
    """Every paragraph, recursing into tables."""
    out = []
    for child in el.iterchildren():
        if child.tag == qn('w:p'):
            out.append(Paragraph(child, parent))
        elif child.tag == qn('w:tbl'):
            for row in Table(child, parent).rows:
                for cell in row.cells:
                    out.extend(cell.paragraphs)
    return out


# 1) Keep lines together on EVERY paragraph (kills one-line widows/orphans).
for p in all_paragraphs(doc.element.body, doc):
    p.paragraph_format.keep_together = True

# 2) Walk the top-level block sequence to glue units together.
body = doc.element.body
blocks = [c for c in body.iterchildren() if c.tag in (qn('w:p'), qn('w:tbl'))]

for i, child in enumerate(blocks):
    nxt = blocks[i + 1] if i + 1 < len(blocks) else None

    if child.tag == qn('w:p'):
        p = Paragraph(child, doc)
        # Glue a paragraph to what follows when the next block is a table
        # (heading/directions -> table) or the next block is a list item, or
        # this is a list item followed by another list item.
        if nxt is not None:
            if nxt.tag == qn('w:tbl'):
                p.paragraph_format.keep_with_next = True
            elif nxt.tag == qn('w:p'):
                # Glue whatever precedes a list item (heading/intro) to it, and
                # chain list items together; the last item's next is non-list,
                # so a page break is still allowed cleanly after the list.
                if has_numpr(Paragraph(nxt, doc)):
                    p.paragraph_format.keep_with_next = True

    elif child.tag == qn('w:tbl'):
        table = Table(child, doc)
        rows = table.rows
        last = len(rows) - 1
        for r_idx, row in enumerate(rows):
            set_cant_split(row)
            if r_idx < last:
                # keep every row with the next -> whole table moves as a unit
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True


doc.save(path)
print("paginated:", path)
